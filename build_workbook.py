import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.properties import PageSetupProperties

"""Builds the Regen Solar PV commissioning workbook (and the cell map the app uses to fill / read it).

    python3 workbook/build_workbook.py
writes workbook/Solar_PV_Commissioning_Report.xlsx, workbook/template_raw.xlsx and workbook/cellmap.json
"""
import json, sys, os, shutil
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(HERE, "Solar_PV_Commissioning_Report.xlsx")
RAW = os.path.join(HERE, "template_raw.xlsx")
MAPOUT = os.path.join(HERE, "cellmap.json")
N_INV = 20          # DC String Test sheets in the workbook (INV 03 … INV 20 start hidden)
N_INV_VISIBLE = 2
TEMPLATE_VERSION = 2
MAP = {"version": TEMPLATE_VERSION}   # cell map used by the app to fill / read this workbook
GREEN = "00984A"
FN = "Arial"
NAVY, BLUE, LBLUE, YEL, GREY, AUTO, WHITE = "1F4E78", "2E75B6", "D9E1F2", "FFF2CC", "F2F2F2", "E7E6E6", "FFFFFF"
side = Side(style="thin", color="808080")
BOX = Border(left=side, right=side, top=side, bottom=side)
FOOT = ("© Regen Renewables (Pvt) Ltd. All rights reserved. This document is confidential and the property of "
        "Regen Renewables (Pvt) Ltd. It must not be copied, shared or used without written permission.")


def fnt(b=False, sz=10, col="000000", it=False):
    return Font(name=FN, size=sz, bold=b, italic=it, color=col)


def fl(c):
    return PatternFill("solid", fgColor=c)


STY = {
    "label": (fnt(True), GREY, "left"),
    "input": (fnt(), YEL, None),
    "calc": (fnt(True), AUTO, "center"),
    "plain": (fnt(), None, None),
    "hdr": (fnt(True, 10, WHITE), BLUE, "center"),
    "dark": (fnt(True, 10, WHITE), NAVY, "center"),
    "sec": (fnt(True, 11), LBLUE, "left"),
}

wb = openpyxl.Workbook()
wb.remove(wb.active)


def new(name, widths, tab):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for k, v in widths.items():
        ws.column_dimensions[k].width = v
    ws.sheet_properties.tabColor = tab
    return ws


def put(ws, r, c1, c2=None, v=None, kind="plain", fmt=None, al=None):
    c2 = c2 or c1
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    font_, fill_, a = STY[kind]
    al = al or a or "center"
    for c in range(c1, c2 + 1):
        x = ws.cell(r, c)
        x.border = BOX
        x.font = font_
        if fill_:
            x.fill = fl(fill_)
        x.alignment = Alignment(horizontal=al, vertical="center", wrap_text=True, indent=1 if al == "left" else 0)
    x = ws.cell(r, c1)
    x.value = v
    if fmt:
        x.number_format = fmt
    return x


def hdr(ws, r, labels, c0=2, h=32):
    ws.row_dimensions[r].height = h
    for i, t in enumerate(labels):
        put(ws, r, c0 + i, None, t, "hdr")


def hdr_spans(ws, r, spans, h=32, kind="hdr"):
    ws.row_dimensions[r].height = h
    for c1, c2, t in spans:
        put(ws, r, c1, c2, t, kind)


def top(ws, last_col, text, size=16, cover=False):
    ws.merge_cells(start_row=2, start_column=2, end_row=3, end_column=last_col)
    for r in (2, 3):
        for c in range(2, last_col + 1):
            x = ws.cell(r, c)
            x.fill = fl(NAVY)
            x.border = BOX
    x = ws.cell(2, 2, text)
    x.font = fnt(True, size, WHITE)
    x.alignment = Alignment(horizontal="center", vertical="center")
    ws.merge_cells(start_row=4, start_column=2, end_row=4, end_column=last_col)
    if cover:
        s = ws.cell(4, 2, "Regen Renewables (Pvt) Ltd.   ·   31/10 Castle Street, Colombo 8, Sri Lanka   ·   regen.lk")
    else:
        s = ws.cell(4, 2, '="Project: "&IF(Cover!$C$7="","(enter on Cover sheet)",Cover!$C$7)&"     |     Regen Renewables (Pvt) Ltd."')
    s.font = fnt(False, 9, "404040", True)
    s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[4].height = 16
    ws.row_dimensions[5].height = 6


def dd(ws, rng, opts):
    dv = DataValidation(type="list", formula1='"%s"' % ",".join(opts), allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(rng)


def cf(ws, rng):
    for txt, bg, fg in (("Pass", "C6EFCE", "006100"), ("Fail", "FFC7CE", "9C0006"),
                        ("Incomplete", "FFEB9C", "9C5700"), ("IN PROGRESS", "FFEB9C", "9C5700"),
                        ("INCOMPLETE", "FFEB9C", "9C5700"), ("N/A", "EDEDED", "595959")):
        ws.conditional_formatting.add(rng, CellIsRule(
            operator="equal", formula=['"%s"' % txt],
            fill=PatternFill(start_color=bg, end_color=bg, fill_type="solid"), font=Font(color=fg, bold=True)))


def note(ws, r, c1, c2, text, h=28):
    ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    x = ws.cell(r, c1, text)
    x.font = fnt(False, 9, "404040", True)
    x.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[r].height = h


def sign_row(ws, r, c_last, split):
    ws.row_dimensions[r].height = 28
    put(ws, r, 2, 3, "Officer Name", "label")
    put(ws, r, 4, split, None, "input", al="left")
    put(ws, r, split + 1, split + 2, "Signature", "label")
    put(ws, r, split + 3, c_last, None, "input", al="left")


def finish(ws, landscape=True, one_page=False):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1 if one_page else 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.6, bottom=0.7, header=0.25, footer=0.3)
    ws.oddHeader.left.text = "Regen Renewables (Pvt) Ltd.  \u2014  Solar PV Installation Commissioning Report"
    ws.oddHeader.left.size = 8
    ws.oddHeader.right.text = "&A   |   Page &P of &N"
    ws.oddHeader.right.size = 8
    ws.oddFooter.center.text = FOOT
    ws.oddFooter.center.size = 7
    ws.print_area = "A1:%s%d" % (CL(ws.max_column), ws.max_row)
    ws.print_options.horizontalCentered = True


def name_cell(name, sheet, ref):
    wb.defined_names[name] = DefinedName(name, attr_text="'%s'!%s" % (sheet, ref))


# =====================================================================
# COVER
# =====================================================================
ws = new("Cover", {"A": 2.5, "B": 32, "C": 24, "D": 3, "E": 32, "F": 24, "G": 2.5}, "595959")
# letterhead: Regen logo on the left, title beside it, green rule underneath
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
LOGO_PNG = os.path.join(ROOT, "assets", "logo", "regen_logo.png")
ws.row_dimensions[1].height = 8
for _r, _h in ((2, 27), (3, 27), (4, 20), (5, 7)):
    ws.row_dimensions[_r].height = _h
ws.merge_cells("B2:B4")
ws.merge_cells("C2:F2"); ws.merge_cells("C3:F3"); ws.merge_cells("C4:F4")
_x = ws["C2"]; _x.value = "SOLAR PV SYSTEM INSTALLATION"; _x.font = fnt(True, 18, NAVY); _x.alignment = Alignment(horizontal="left", vertical="bottom")
_x = ws["C3"]; _x.value = "COMMISSIONING REPORT"; _x.font = fnt(True, 18, NAVY); _x.alignment = Alignment(horizontal="left", vertical="center")
_x = ws["C4"]; _x.value = "Regen Renewables (Pvt) Ltd.   ·   31/10 Castle Street, Colombo 8, Sri Lanka   ·   regen.lk"; _x.font = fnt(False, 9.5, GREEN, True); _x.alignment = Alignment(horizontal="left", vertical="top")
for _c in range(2, 7):
    ws.cell(5, _c).border = Border(bottom=Side(style="thick", color=GREEN))
_lg = XLImage(LOGO_PNG)
_h = 92; _w = round(_h * _lg.width / _lg.height)
_lg.width, _lg.height = _w, _h
_EMU = 9525
_lg.anchor = OneCellAnchor(_from=AnchorMarker(col=1, colOff=4 * _EMU, row=1, rowOff=2 * _EMU), ext=XDRPositiveSize2D(_w * _EMU, _h * _EMU))
ws.add_image(_lg)


def pair_block(r, title_text, left, right, dds=None):
    put(ws, r, 2, 6, title_text, "sec")
    ws.row_dimensions[r].height = 20
    r += 1
    pos = {}
    for (a, va), (b, vb) in zip(left, right):
        ws.row_dimensions[r].height = 22
        put(ws, r, 2, None, a, "label")
        put(ws, r, 3, None, va, "input", al="left")
        put(ws, r, 5, None, b, "label")
        put(ws, r, 6, None, vb, "input", al="left")
        pos[a] = ("C", r)
        pos[b] = ("F", r)
        r += 1
    return r, pos


r, pos = pair_block(6, "1.  GENERAL DETAILS",
                    [("Project Name", None), ("Address", None), ("Contact Person Name", None),
                     ("Telephone — Fixed", None), ("E-Mail", None), ("Electricity Bill Available", None)],
                    [("Document No.", None), ("Revision No.", "Rev 0"), ("Commissioning Date", None),
                     ("Telephone — Mobile", None), ("Weather Condition", None), ("Client", None)])
assert pos["Project Name"] == ("C", 7) and pos["Commissioning Date"] == ("F", 9)
COVER_POS = dict(pos)
dd(ws, "C12:C12", ["Yes", "No"])

r += 1
r2, pos2 = pair_block(r, "2.  PV SYSTEM & UTILITY DETAILS",
                      [("EPC / Installation Contractor", None), ("PV System Capacity (kWp)", None),
                       ("System AC Capacity (kW)", None), ("PV System Connection Method", None),
                       ("Utility Supplier", None), ("Utility Account No.", None)],
                      [("Inverter Brand & Model", None), ("Number of Inverters", None),
                       ("Roof Structure", None), ("Utility System Voltage", None),
                       ("Utility Supplier's Branch Office", None), ("Installation Type", None)])
def cref(col, key):
    rr = pos2[key][1]
    return "%s%d:%s%d" % (col, rr, col, rr)


dd(ws, cref("C", "PV System Connection Method"), ["Net Metering", "Net Accounting", "Net Plus"])
dd(ws, cref("C", "Utility Supplier"), ["CEB", "LECO"])
dd(ws, cref("F", "Utility System Voltage"), ["230V", "400V"])
dd(ws, cref("F", "Installation Type"), ["Grid-tied", "Hybrid (with battery)", "Off-grid"])

COVER_POS.update(pos2)
r = r2 + 1
r3, _p3 = pair_block(r, "3.  PV PANEL DETAILS",
                   [("PV Module Brand & Model", None), ("Voc / Vmp (V)", None), ("Number of Panels", None)],
                   [("Rated Power (Wp)", None), ("Imp / Isc (A)", None), ("Number of Strings", None)])
COVER_POS.update(_p3)
r = r3 + 1
r4, _p4 = pair_block(r, "4.  PREPARED / REVIEWED / APPROVED",
                   [("Prepared By (Name / Company)", None), ("Approved By (Name / Company)", None)],
                   [("Reviewed By (Name / Company)", None), ("Approval Date", None)])
COVER_POS.update(_p4)
_ck = {"Project Name": "name", "Address": "address", "Contact Person Name": "contact", "Telephone — Fixed": "tel_fixed",
       "E-Mail": "email", "Electricity Bill Available": "bill", "Document No.": "docno", "Revision No.": "rev",
       "Commissioning Date": "date", "Telephone — Mobile": "tel_mobile", "Weather Condition": "weather", "Client": "client",
       "EPC / Installation Contractor": "epc", "PV System Capacity (kWp)": "kwp", "System AC Capacity (kW)": "kwac",
       "PV System Connection Method": "conn", "Utility Supplier": "supplier", "Utility Account No.": "account",
       "Inverter Brand & Model": "inv_model", "Number of Inverters": "ninv", "Roof Structure": "roof",
       "Utility System Voltage": "voltage", "Utility Supplier's Branch Office": "branch", "Installation Type": "itype",
       "PV Module Brand & Model": "pv_model", "Voc / Vmp (V)": "pv_vv", "Number of Panels": "npan",
       "Rated Power (Wp)": "pv_wp", "Imp / Isc (A)": "pv_ii", "Number of Strings": "nstr",
       "Prepared By (Name / Company)": "prep", "Approved By (Name / Company)": "appr",
       "Reviewed By (Name / Company)": "rev_by", "Approval Date": "appr_date"}
MAP["cover"] = {k: "%s%d" % COVER_POS[lab] for lab, k in _ck.items()}
for _k in ("date", "appr_date"):
    ws[MAP["cover"][_k]].number_format = "dd-mmm-yyyy"
r = r4 + 1
put(ws, r, 2, 6, "5.  HOW THIS WORKBOOK WORKS", "sec")
ws.row_dimensions[r].height = 20
r += 1
for t in [
    "Yellow cells = fill in on site.  Grey cells = calculated automatically (Pass / Fail, totals, deviations) — do not type over them.",
    "Every reading is entered ONCE, in its own table. The Summary sheet collects all results automatically, and the project name is picked up from this Cover on every sheet.",
    "Acceptance limits (minimum insulation, earth resistance, voltage / frequency window, thermal ΔT) are on “Equipment and Limits”. They are defaults — edit them to your company / utility standard.",
    "Up to 20 inverters: sheets “DC INV 03” … “DC INV 20” are hidden — right-click any sheet tab → Unhide, then on the Summary sheet unhide that inverter’s row and set Applicable to Yes. MPPT / String columns are free-text, so any layout works.",
    "Insulation testers that show “>” or “OL”: type the highest number displayed (e.g. 999). Single-phase inverters: type N/A in the L2 / L3 insulation cells.",
    "Battery Test applies to hybrid systems only — set “Applicable” to Yes on the Summary sheet when a battery is installed.",
]:
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=6)
    x = ws.cell(r, 2, "•  " + t)
    x.font = fnt(False, 9.5)
    x.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
    ws.row_dimensions[r].height = 30
    r += 1
r += 1
put(ws, r, 2, 6, "6.  REPORT CONTENTS", "sec")
ws.row_dimensions[r].height = 20
r += 1
contents = ["Summary (auto)", "Equipment and Limits", "Visual Inspection", "Inverters and Panel",
            "DC String Tests — one sheet per inverter (up to 20)", "Panel Serial Numbers", "AC Output (incl. grid no-load)",
            "AC Insulation and Earthing", "Thermal Imaging", "Battery (hybrid only)",
            "Protection and Grid", "Punch List", "Sign-Off & Handover"]
for i, t in enumerate(contents, 1):
    ws.row_dimensions[r].height = 18
    put(ws, r, 2, None, i, "plain")
    put(ws, r, 3, 6, t, "plain", al="left")
    r += 1
finish(ws, landscape=False, one_page=True)

# =====================================================================
# EQUIPMENT AND LIMITS
# =====================================================================
EQ = "Equipment and Limits"
ws = new(EQ, {"A": 2.5, "B": 6, "C": 38, "D": 26, "E": 18, "F": 18, "G": 18, "H": 36}, "595959")
top(ws, 8, "TEST EQUIPMENT & ACCEPTANCE LIMITS")
put(ws, 6, 2, 8, "TEST EQUIPMENT REGISTER  —  recorded once, applies to every test in this report", "sec")
hdr(ws, 7, ["No.", "Instrument", "Model / Make", "Serial No.", "Calibrated On", "Calibration Due", "Used For"])
inst = [("Insulation Resistance Tester", None, "DC & AC insulation tests"),
        ("Digital Multimeter / Clamp Meter", None, "Voltage, current, polarity"),
        ("PV String Tester / I-V Curve Tracer", None, "Voc, Isc, string current"),
        ("Digital Earth Tester", "4105A / KYORITSU", "Earthing test"),
        ("Thermal Imaging Camera", None, "Panels, inverters, switchgear"),
        ("Power / Energy Analyser", None, "AC output, PF, frequency"),
        ("Irradiance Meter", None, "Irradiance during on-load tests"),
        ("Other", None, None)]
r = 8
for i, (a, m, u) in enumerate(inst, 1):
    ws.row_dimensions[r].height = 22
    put(ws, r, 2, None, i)
    put(ws, r, 3, None, a, "plain", al="left")
    put(ws, r, 4, None, m, "input", al="left")
    put(ws, r, 5, None, None, "input")
    put(ws, r, 6, None, None, "input", fmt="dd-mmm-yyyy")
    put(ws, r, 7, None, None, "input", fmt="dd-mmm-yyyy")
    put(ws, r, 8, None, u, "input", al="left")
    r += 1
MAP["equip"] = {"sheet": EQ, "first": 8, "n": len(inst), "cols": {"model": "D", "sn": "E", "cal": "F", "due": "G", "use": "H"}}
r += 1
put(ws, r, 2, 8, "ACCEPTANCE LIMITS  —  defaults, edit to match your company / utility standard", "sec")
r += 1
hdr(ws, r, ["No.", "Parameter", "Value", "Unit"], h=24)
put(ws, r, 6, 8, "Used in", "hdr")
r += 1
limits = [
    ("LIM_INS_DC", "DC insulation resistance — minimum", 1, "MΩ", "DC String Tests", None),
    ("LIM_INS_AC", "AC insulation resistance — minimum", 1, "MΩ", "AC Insulation and Earthing", None),
    ("LIM_VOC_DEV", "String Voc deviation within same MPPT / module count — maximum", 0.05, "%", "DC String Tests", "0%"),
    ("LIM_EARTH", "Earth resistance — maximum", 5, "Ω", "AC Insulation and Earthing", None),
    ("LIM_VNOM", "Nominal phase voltage (L–N)", 230, "V", "AC Output", None),
    ("LIM_VTOL", "Voltage tolerance (±)", 0.10, "%", "AC Output", "0%"),
    ("LIM_FNOM", "Nominal frequency", 50, "Hz", "AC Output", None),
    ("LIM_FTOL", "Frequency tolerance (±)", 1, "Hz", "AC Output", None),
    ("LIM_DT", "Thermal ΔT above ambient — maximum", 20, "°C", "Thermal Imaging", None),
]
for i, (nm, par, val, unit, used, fmt) in enumerate(limits, 1):
    ws.row_dimensions[r].height = 22
    put(ws, r, 2, None, i)
    put(ws, r, 3, None, par, "plain", al="left")
    put(ws, r, 4, None, val, "input", fmt=fmt)
    put(ws, r, 5, None, unit)
    put(ws, r, 6, 8, used, "plain", al="left")
    name_cell(nm, EQ, "$D$%d" % r)
    MAP.setdefault("limits", {})[{"LIM_INS_DC": "insdc", "LIM_INS_AC": "insac", "LIM_VOC_DEV": "vocdev", "LIM_EARTH": "earth", "LIM_VNOM": "vnom", "LIM_VTOL": "vtol", "LIM_FNOM": "fnom", "LIM_FTOL": "ftol", "LIM_DT": "dt"}[nm]] = {"cell": "D%d" % r, "pct": fmt == "0%"}
    r += 1
finish(ws)

# =====================================================================
# VISUAL INSPECTION
# =====================================================================
INS = "Visual Inspection"
ws = new(INS, {"A": 2.5, "B": 6, "C": 58, "D": 14, "E": 38}, "0E7C7B")
top(ws, 5, "VISUAL & MECHANICAL INSPECTION CHECKLIST", 15)
hdr(ws, 6, ["No.", "Inspection Item", "Status", "Remarks"], h=26)
groups = {
    "A.  PV Modules & Mounting Structure": [
        "Modules free of physical damage (cracks, hot spots, delamination)",
        "Modules clean and free of shading obstructions",
        "Mounting structure secure, level and correctly torqued",
        "Module clamps / bolts torqued to spec, no loose fixings",
        "Structure and module frames bonded to earthing system"],
    "B.  DC Cabling": [
        "DC cables routed and dressed neatly, no sharp bends",
        "Cables secured with UV-rated cable ties / clips",
        "MC4 (or equivalent) connectors fully locked, no exposed conductor",
        "String / cable labelling complete and legible",
        "Cable entry points and conduits sealed (IP rated)"],
    "C.  Inverter & DC Protection": [
        "Inverter securely mounted, correct orientation and clearance",
        "Ventilation / cooling clearance around inverter as per manual",
        "Cable glands sealed, unused entries blanked off",
        "DC surge protection device (SPD) fitted and indicator normal",
        "DC isolator / switch labelled and operates freely",
        "Inverter display / status LEDs functioning normally"],
    "D.  AC Side & Main Panel": [
        "AC cable terminations tight and correctly torqued",
        "AC circuit breaker rating matches design / datasheet",
        "AC isolator accessible and labelled",
        "AC cable and circuit labelling complete",
        "Energy meter / CT connections correct and secure"],
    "E.  Earthing & Lightning Protection": [
        "Earth pit(s) installed and accessible for testing",
        "Earth cable size / type as per design",
        "Equipotential bonding complete across all metallic parts",
        "Lightning / surge protection installed where specified"],
    "F.  Safety & Signage": [
        "Danger / warning labels posted at DC & AC points",
        "Single line diagram (SLD) posted at site / inverter room",
        "Fire extinguisher (CO2 / DCP, rated for electrical) available",
        "PPE available for O&M personnel",
        "Emergency shutdown procedure posted"],
}
r = 7
n = 1
INS_FIRST = r
for g, items in groups.items():
    put(ws, r, 2, 5, g, "sec")
    ws.row_dimensions[r].height = 20
    r += 1
    for it in items:
        ws.row_dimensions[r].height = 22
        put(ws, r, 2, None, n)
        put(ws, r, 3, None, it, "plain", al="left")
        put(ws, r, 4, None, None, "input")
        put(ws, r, 5, None, None, "input", al="left")
        MAP.setdefault("visual", {"sheet": INS, "rows": []})["rows"].append(r)
        n += 1
        r += 1
INS_LAST = r - 1
dd(ws, "D%d:D%d" % (INS_FIRST, INS_LAST), ["Pass", "Fail", "N/A"])
cf(ws, "D%d:D%d" % (INS_FIRST, INS_LAST))
r += 1
put(ws, r, 2, 3, "Inspected By", "label")
put(ws, r, 4, 5, None, "input", al="left")
MAP["visual"]["by"] = "D%d" % r
r += 1
put(ws, r, 2, 3, "Date", "label")
put(ws, r, 4, 5, None, "input", al="left", fmt="dd-mmm-yyyy")
MAP["visual"]["date"] = "D%d" % r
finish(ws, landscape=False)

# =====================================================================
# INVERTERS AND PANEL
# =====================================================================
INV = "Inverters and Panel"
ws = new(INV, {"A": 2.5, "B": 5, "C": 17, "D": 13, "E": 15, "F": 15, "G": 10, "H": 9, "I": 14,
               "J": 9, "K": 11, "L": 15, "M": 10, "N": 11, "O": 22}, "1F4E78")
top(ws, 15, "INVERTER REGISTER & MAIN PANEL DETAILS")
put(ws, 6, 2, 15, "INVERTER REGISTER  —  one row per inverter (strings / PV modules per inverter are totalled on each DC String Test sheet)", "sec")
hdr(ws, 7, ["No.", "Location / Room", "Make", "Model", "Serial No.", "Rated Power (kW)", "DC SPD (Y/N)",
            "DC SPD Model", "DC Isolator Qty", "DC Isolator Rating", "DC Cable (size / type)",
            "AC Isolator (Y/N)", "AC Isolator Rating", "Remarks"], h=44)
for i in range(1, N_INV + 1):
    r = 7 + i
    ws.row_dimensions[r].height = 22
    put(ws, r, 2, None, i)
    for c in range(3, 16):
        put(ws, r, c, None, None, "input", al="left" if c in (3, 4, 5, 6, 9, 12, 15) else None)
RG_LAST = 7 + N_INV
dd(ws, "H8:H%d" % RG_LAST, ["Y", "N"])
dd(ws, "M8:M%d" % RG_LAST, ["Y", "N"])
MAP["invreg"] = {"sheet": INV, "first": 8, "n": N_INV, "cols": {"loc": "C", "make": "D", "model": "E", "sn": "F", "kw": "G", "spd": "H", "spd_model": "I", "iso_qty": "J", "iso_rating": "K", "dc_cable": "L", "ac_iso": "M", "ac_iso_rating": "N", "rem": "O"}}
PB = RG_LAST + 2
put(ws, PB, 2, 15, "MAIN ELECTRICAL PANEL & CABLE DETAILS", "sec")
left = ["Main panel location", "Main breaker make / model", "Main breaker rating (A)", "Enclosure IP rating",
        "AC SPD fitted (Y/N)", "Energy meter make / model"]
right = ["Energy meter serial no.", "CT ratio (if used)", "RCD / ELCB rating (mA)",
         "AC cable: inverter → panel (size / type / length)", "AC cable: panel → utility (size / type / length)",
         "Earth conductor size"]
lk = ["loc", "brk", "brk_a", "ip", "ac_spd", "meter"]; rk = ["meter_sn", "ct", "rcd", "cable_inv", "cable_util", "earth"]
MAP["panel"] = {"sheet": INV}
r = PB + 1
for a, b, ka, kb in zip(left, right, lk, rk):
    ws.row_dimensions[r].height = 30
    put(ws, r, 2, 4, a, "label")
    put(ws, r, 5, 7, None, "input", al="left")
    put(ws, r, 8, 11, b, "label")
    put(ws, r, 12, 15, None, "input", al="left")
    MAP["panel"][ka] = "E%d" % r; MAP["panel"][kb] = "L%d" % r
    r += 1
dd(ws, MAP["panel"]["ac_spd"] + ":" + MAP["panel"]["ac_spd"], ["Y", "N"])
finish(ws)

# =====================================================================
# DC STRING TEST (one sheet per inverter)
# =====================================================================
DC_FIRST, DC_LAST = 15, 50
DCS = {}


def dc_sheet(inv):
    nm = "DC " + inv
    ws = new(nm, {"A": 2.5, "B": 5, "C": 9, "D": 9, "E": 9, "F": 10, "G": 10, "H": 11, "I": 11, "J": 11,
                  "K": 10, "L": 11, "M": 11, "N": 10, "O": 26}, "1F4E78")
    top(ws, 15, "DC STRING TEST — " + inv)
    put(ws, 6, 2, 15, "INVERTER DETAILS", "sec")
    rows = [("Inverter No.", "Inverter Location / Room"), ("Make / Model", "Serial No."),
            ("Rated Capacity (kW)", "Date / Time"), ("Tested By", "Weather / Irradiance (W/m²)")]
    r = 7
    for a, b in rows:
        ws.row_dimensions[r].height = 22
        put(ws, r, 2, 5, a, "label")
        put(ws, r, 6, 8, inv if a == "Inverter No." else None, "input", al="left")
        put(ws, r, 9, 11, b, "label")
        put(ws, r, 12, 15, None, "input", al="left")
        r += 1
    ws.row_dimensions[r].height = 22
    put(ws, r, 2, 5, "Total Strings Listed", "label")
    put(ws, r, 6, 8, "=COUNT(E%d:E%d)" % (DC_FIRST, DC_LAST), "calc")
    put(ws, r, 9, 11, "Total PV Modules", "label")
    put(ws, r, 12, 15, "=SUM(E%d:E%d)" % (DC_FIRST, DC_LAST), "calc")
    put(ws, 12, 2, 15, "STRING TESTS  —  each string is entered ONCE (MPPT / string / module count are not repeated)", "sec")
    hdr_spans(ws, 13, [(2, 5, "STRING DETAILS"), (6, 11, "NO-LOAD  (before energising)"),
                       (12, 13, "ON-LOAD"), (14, 15, "RESULT")], h=20, kind="dark")
    hdr(ws, 14, ["No.", "MPPT", "String", "PV Modules", "Voc (V)", "Voc dev. vs group",
                 "Insulation + to Earth (MΩ)", "Insulation – to Earth (MΩ)", "Insulation + to – (MΩ)",
                 "Polarity", "Operating Voltage (V)", "String Current (A)", "Result", "Remarks"], h=46)
    F, L = DC_FIRST, DC_LAST
    for i, r in enumerate(range(F, L + 1), 1):
        ws.row_dimensions[r].height = 19
        put(ws, r, 2, None, i)
        put(ws, r, 3, None, None, "input")
        put(ws, r, 4, None, None, "input")
        put(ws, r, 5, None, None, "input", fmt="0")
        put(ws, r, 6, None, None, "input", fmt="0.0")
        put(ws, r, 7, None,
            '=IF(NOT(ISNUMBER(F{r})),"",IFERROR(F{r}/IF(AND(C{r}<>"",E{r}<>""),AVERAGEIFS(F${F}:F${L},C${F}:C${L},C{r},E${F}:E${L},E{r}),AVERAGE(F${F}:F${L}))-1,""))'.format(r=r, F=F, L=L),
            "calc", fmt="0.0%")
        for c in (8, 9, 10):
            put(ws, r, c, None, None, "input", fmt="0.0")
        put(ws, r, 11, None, None, "input")
        put(ws, r, 12, None, None, "input", fmt="0.0")
        put(ws, r, 13, None, None, "input", fmt="0.00")
        put(ws, r, 14, None,
            '=IF(COUNTA(F{r},H{r}:M{r})=0,"",IF(OR(K{r}="Reverse",AND(COUNT(H{r}:J{r})>0,MIN(H{r}:J{r})<LIM_INS_DC),IF(ISNUMBER(G{r}),ABS(G{r})>LIM_VOC_DEV,FALSE)),"Fail",IF(AND(ISNUMBER(F{r}),COUNT(H{r}:J{r})=3,K{r}="OK"),"Pass","Incomplete")))'.format(r=r),
            "calc")
        put(ws, r, 15, None, None, "input", al="left")
    dd(ws, "K%d:K%d" % (F, L), ["OK", "Reverse"])
    cf(ws, "N%d:N%d" % (F, L))
    r = L + 2
    for lab_, key in (("Strings — Pass", "Pass"), ("Strings — Fail", "Fail")):
        ws.row_dimensions[r].height = 22
        put(ws, r, 2, 5, lab_, "label")
        put(ws, r, 6, 8, '=COUNTIF(N%d:N%d,"%s")' % (F, L, key), "calc")
        r += 1
    ws.row_dimensions[r - 2].height = 22
    put(ws, r - 2, 9, 11, "Strings — Incomplete", "label")
    put(ws, r - 2, 12, 15, '=COUNTIF(N%d:N%d,"Incomplete")' % (F, L), "calc")
    r += 1
    sign_row(ws, r, 15, 8)
    MAP["dc"]["officer"] = "D%d" % r
    ws.print_title_rows = "13:14"
    finish(ws)
    DCS[inv] = nm
    return ws


MAP["dc"] = {"first": DC_FIRST, "last": DC_LAST, "head": {"loc": "L7", "makemodel": "F8", "sn": "L8", "kw": "F9", "dt": "L9", "by": "F10", "wx": "L10"},
             "cols": {"mppt": "C", "str": "D", "mods": "E", "voc": "F", "ip": "H", "in": "I", "ipn": "J", "pol": "K", "opv": "L", "cur": "M", "rem": "O"},
             "num": ["mods", "voc", "ip", "in", "ipn", "opv", "cur"], "sheets": []}
for _k in range(1, N_INV + 1):
    _ws = dc_sheet("INV %02d" % _k)
    MAP["dc"]["sheets"].append(_ws.title)
    if _k > N_INV_VISIBLE:
        _ws.sheet_state = "hidden"

# =====================================================================
# AC OUTPUT
# =====================================================================
ACO = "AC Output"
ws = new(ACO, {"A": 2.5, "B": 5, "C": 22, "D": 9, "E": 10, "F": 10, "G": 10, "H": 10, "I": 10, "J": 10,
               "K": 10, "L": 8, "M": 10, "N": 10, "O": 24}, "C55A11")
top(ws, 15, "AC OUTPUT TEST  (grid no-load and inverter on-load)")
put(ws, 6, 2, 15, "AC SIDE READINGS  —  first row = grid side with inverters OFF, then each inverter ON-LOAD", "sec")
ws.row_dimensions[7].height = 22
put(ws, 7, 2, 3, "Date", "label")
put(ws, 7, 4, 7, None, "input", fmt="dd-mmm-yyyy")
put(ws, 7, 8, 10, "Tested By", "label")
put(ws, 7, 11, 15, None, "input", al="left")
ws.row_dimensions[8].height = 6
hdr_spans(ws, 9, [(2, 4, "POINT"), (5, 7, "VOLTAGE L–N (V)"), (8, 10, "CURRENT (A)"),
                  (11, 13, "POWER"), (14, 15, "RESULT")], h=20, kind="dark")
hdr(ws, 10, ["No.", "Grid / Inverter & Room", "Time", "L1", "L2", "L3", "L1", "L2", "L3",
             "Active Power (kW)", "PF", "Freq (Hz)", "Result", "Remarks"], h=34)
A_FIRST, A_LAST = 11, 32
for i, r in enumerate(range(A_FIRST, A_LAST + 1), 1):
    ws.row_dimensions[r].height = 20
    put(ws, r, 2, None, i)
    put(ws, r, 3, None, "Grid side (no load)" if i == 1 else None, "input", al="left")
    put(ws, r, 4, None, None, "input", fmt="hh:mm")
    for c in range(5, 8):
        put(ws, r, c, None, None, "input", fmt="0.0")
    for c in range(8, 11):
        put(ws, r, c, None, None, "input", fmt="0.0")
    put(ws, r, 11, None, None, "input", fmt="0.00")
    put(ws, r, 12, None, None, "input", fmt="0.00")
    put(ws, r, 13, None, None, "input", fmt="0.00")
    put(ws, r, 14, None,
        '=IF(COUNT(E{r}:G{r})=0,"",IF(OR(MAX(E{r}:G{r})>LIM_VNOM*(1+LIM_VTOL),MIN(E{r}:G{r})<LIM_VNOM*(1-LIM_VTOL),IF(ISNUMBER(M{r}),ABS(M{r}-LIM_FNOM)>LIM_FTOL,FALSE)),"Fail","Pass"))'.format(r=r),
        "calc")
    put(ws, r, 15, None, None, "input", al="left")
cf(ws, "N%d:N%d" % (A_FIRST, A_LAST))
ws.row_dimensions[A_LAST + 2].height = 22
put(ws, A_LAST + 2, 2, 5, "Total plant AC power (kW)", "label")
put(ws, A_LAST + 2, 6, 8, "=SUM(K%d:K%d)" % (A_FIRST, A_LAST), "calc", fmt="0.00")
note(ws, A_LAST + 3, 2, 15, "Result checks every phase against the voltage window and the frequency against the tolerance set on “Equipment and Limits”. Leave unused phases blank for single-phase inverters.", 24)
sign_row(ws, A_LAST + 5, 15, 8)
MAP["ac"] = {"sheet": ACO, "date": "D7", "by": "K7", "first": A_FIRST, "last": A_LAST, "officer": "D%d" % (A_LAST + 5),
             "cols": {"pt": "C", "time": "D", "v1": "E", "v2": "F", "v3": "G", "i1": "H", "i2": "I", "i3": "J", "kw": "K", "pf": "L", "hz": "M", "rem": "O"},
             "num": ["v1", "v2", "v3", "i1", "i2", "i3", "kw", "pf", "hz"], "time": ["time"]}
finish(ws)

# =====================================================================
# AC INSULATION AND EARTHING
# =====================================================================
ACI = "AC Insulation and Earthing"
ws = new(ACI, {"A": 2.5, "B": 6, "C": 24, "D": 9, "E": 9, "F": 9, "G": 9, "H": 9, "I": 9, "J": 9,
               "K": 9, "L": 9, "M": 9, "N": 10, "O": 24}, "C55A11")
top(ws, 15, "AC INSULATION & EARTHING TEST")
put(ws, 6, 2, 15, "AC SIDE INSULATION RESISTANCE (MΩ)  —  one row per inverter / circuit (grid side and main panel can be added as extra rows)", "sec")
ws.row_dimensions[7].height = 22
put(ws, 7, 2, 3, "Date", "label")
put(ws, 7, 4, 6, None, "input", fmt="dd-mmm-yyyy")
put(ws, 7, 7, 9, "Test Voltage (V)", "label")
put(ws, 7, 10, 11, None, "input")
put(ws, 7, 12, 13, "Tested By", "label")
put(ws, 7, 14, 15, None, "input", al="left")
ws.row_dimensions[8].height = 6
hdr(ws, 9, ["No.", "Inverter / Circuit", "L1–L2", "L1–L3", "L2–L3", "L1–N", "L2–N", "L3–N",
            "L1–E", "L2–E", "L3–E", "N–E", "Result", "Remarks"], h=28)
I_FIRST, I_LAST = 10, 30
for i, r in enumerate(range(I_FIRST, I_LAST + 1), 1):
    ws.row_dimensions[r].height = 20
    put(ws, r, 2, None, i)
    put(ws, r, 3, None, None, "input", al="left")
    for c in range(4, 14):
        put(ws, r, c, None, None, "input", fmt="0.0")
    put(ws, r, 14, None,
        '=IF(COUNTA(D{r}:M{r})=0,"",IF(AND(COUNT(D{r}:M{r})>0,MIN(D{r}:M{r})<LIM_INS_AC),"Fail",IF(COUNTA(D{r}:M{r})=10,"Pass","Incomplete")))'.format(r=r),
        "calc")
    put(ws, r, 15, None, None, "input", al="left")
cf(ws, "N%d:N%d" % (I_FIRST, I_LAST))
note(ws, I_LAST + 1, 2, 15, "Single-phase: type N/A in the cells that do not apply. If the tester shows “>” or “OL”, type the highest number displayed.", 22)
r = I_LAST + 3
put(ws, r, 2, 15, "EARTHING TEST", "sec")
r += 1
hdr_spans(ws, r, [(2, 2, "Rod No."), (3, 3, "Location"), (4, 6, "Value (Ω)"), (7, 8, "Result"), (9, 15, "Remarks")], h=28)
E_FIRST = r + 1
for i in range(10):
    r += 1
    ws.row_dimensions[r].height = 20
    put(ws, r, 2, None, i + 1)
    put(ws, r, 3, None, None, "input", al="left")
    put(ws, r, 4, 6, None, "input", fmt="0.00")
    put(ws, r, 7, 8, '=IF(NOT(ISNUMBER(D{r})),"",IF(D{r}<=LIM_EARTH,"Pass","Fail"))'.format(r=r), "calc")
    put(ws, r, 9, 15, None, "input", al="left")
E_LAST = r
cf(ws, "G%d:G%d" % (E_FIRST, E_LAST))
r += 2
sign_row(ws, r, 15, 8)
MAP["acins"] = {"sheet": ACI, "date": "D7", "volt": "J7", "by": "N7", "first": I_FIRST, "last": I_LAST, "officer": "D%d" % r,
                "cols": {"pt": "C", "rem": "O"}, "vcols": list("DEFGHIJKLM"),
                "earth": {"first": E_FIRST, "last": E_LAST, "cols": {"loc": "C", "val": "D", "rem": "I"}, "num": ["val"]}}
finish(ws)

# =====================================================================
# THERMAL IMAGING
# =====================================================================
TH = "Thermal Imaging"
ws = new(TH, {"A": 2.5, "B": 6, "C": 26, "D": 20, "E": 10, "F": 12, "G": 12, "H": 11, "I": 34}, "C55A11")
top(ws, 9, "THERMAL IMAGING TEST")
put(ws, 6, 2, 9, "TEST CONDITIONS", "sec")
ws.row_dimensions[7].height = 22
put(ws, 7, 2, 3, "Date", "label")
put(ws, 7, 4, 5, None, "input", fmt="dd-mmm-yyyy")
put(ws, 7, 6, 7, "Ambient Temp (°C)", "label")
put(ws, 7, 8, 9, None, "input", fmt="0.0")
ws.row_dimensions[8].height = 22
put(ws, 8, 2, 3, "Irradiance (W/m²)", "label")
put(ws, 8, 4, 5, None, "input")
put(ws, 8, 6, 7, "Tested By", "label")
put(ws, 8, 8, 9, None, "input", al="left")
AMB = "$H$7"
TH_RES = []
TH_ROWS = []


def th_table(r, title_, labels, rows, kind, prefill=None):
    put(ws, r, 2, 9, title_, "sec")
    r += 1
    hdr(ws, r, labels, h=34)
    first = r + 1
    for i in range(rows):
        r += 1
        ws.row_dimensions[r].height = 20
        put(ws, r, 2, None, i + 1)
        put(ws, r, 3, None, prefill if (prefill and i == 0) else None, "input", al="left")
        put(ws, r, 4, None, None, "input", al="left")
        put(ws, r, 5, None, None, "input", fmt="hh:mm")
        put(ws, r, 6, None, None, "input", fmt="0.0")
        if kind == "single":
            put(ws, r, 7, None, '=IF(AND(ISNUMBER(F{r}),ISNUMBER({a})),F{r}-{a},"")'.format(r=r, a=AMB), "calc", fmt="0.0")
            put(ws, r, 8, None,
                '=IF(NOT(ISNUMBER(F{r})),"",IF(NOT(ISNUMBER(G{r})),"Incomplete",IF(G{r}>LIM_DT,"Fail","Pass")))'.format(r=r), "calc")
        else:
            put(ws, r, 7, None, None, "input", fmt="0.0")
            put(ws, r, 8, None,
                '=IF(COUNT(F{r}:G{r})=0,"",IF(NOT(ISNUMBER({a})),"Incomplete",IF(MAX(F{r}:G{r})-{a}>LIM_DT,"Fail","Pass")))'.format(r=r, a=AMB), "calc")
        put(ws, r, 9, None, None, "input", al="left")
    cf(ws, "H%d:H%d" % (first, r))
    TH_RES.append("$H$%d:$H$%d" % (first, r))
    TH_ROWS.append((first, r))
    return r + 2


r = 10
r = th_table(r, "SOLAR PANELS — CRITICAL / HOT PANELS",
             ["No.", "Inverter / String", "Panel Serial No.", "Time", "Temp (°C)", "ΔT vs ambient (°C)", "Result", "Comment"],
             12, "single")
r = th_table(r, "INVERTERS",
             ["No.", "Inverter / Room", "Point Inspected", "Time", "Temp (°C)", "ΔT vs ambient (°C)", "Result", "Comment"],
             20, "single")
r = th_table(r, "SWITCHGEAR PANEL — BREAKERS",
             ["No.", "Breaker (Main / Inverter)", "Rating / Ref", "Time", "IN Temp (°C)", "OUT Temp (°C)", "Result", "Comment"],
             22, "double", prefill="Main Breaker")
note(ws, r - 1, 2, 9, "Result = Fail when the hottest reading is more than the ΔT limit (see “Equipment and Limits”) above ambient. Attach thermal images to the report separately.", 24)
sign_row(ws, r + 1, 9, 4)
MAP["th"] = {"sheet": TH, "date": "D7", "amb": "H7", "irr": "D8", "by": "H8", "officer": "D%d" % (r + 1),
             "pan": {"first": TH_ROWS[0][0], "last": TH_ROWS[0][1], "cols": {"where": "C", "sn": "D", "time": "E", "t": "F", "c": "I"}, "num": ["t"], "time": ["time"]},
             "inv": {"first": TH_ROWS[1][0], "last": TH_ROWS[1][1], "cols": {"where": "C", "pt": "D", "time": "E", "t": "F", "c": "I"}, "num": ["t"], "time": ["time"]},
             "sw": {"first": TH_ROWS[2][0], "last": TH_ROWS[2][1], "cols": {"where": "C", "rating": "D", "time": "E", "tin": "F", "tout": "G", "c": "I"}, "num": ["tin", "tout"], "time": ["time"]}}
finish(ws)

# =====================================================================
# BATTERY
# =====================================================================
BAT = "Battery"
ws = new(BAT, {"A": 2.5, "B": 6, "C": 28, "D": 16, "E": 16, "F": 16, "G": 16, "H": 14, "I": 14, "J": 24}, "1F6E43")
top(ws, 10, "BATTERY TEST  (hybrid / battery-backed systems)")
put(ws, 6, 2, 10, "BATTERY SYSTEM DETAILS", "sec")
left = ["Battery Brand & Model", "Battery Type", "Bank Capacity (kWh)", "Nominal Voltage (V)"]
right = ["No. of Modules / Racks", "No. of Battery Strings", "BMS Model", "Date / Tested By"]
r = 7
for a, b in zip(left, right):
    ws.row_dimensions[r].height = 22
    put(ws, r, 2, 3, a, "label")
    put(ws, r, 4, 5, None, "input", al="left")
    put(ws, r, 6, 7, b, "label")
    put(ws, r, 8, 10, None, "input", al="left")
    r += 1
dd(ws, "D8:D8", ["Lithium-ion", "Lead-Acid", "Other"])
r += 1
put(ws, r, 2, 10, "BATTERY MODULE / STRING VOLTAGE & INSULATION", "sec")
r += 1
hdr_spans(ws, r, [(2, 2, "No."), (3, 3, "Module / Rack No."), (4, 4, "String No."), (5, 5, "Nominal Voltage (V)"),
                  (6, 6, "Measured Voltage (V)"), (7, 7, "Insulation (MΩ)"), (8, 8, "Status"), (9, 10, "Remarks")], h=34)
BM_FIRST = r + 1
for i in range(24):
    r += 1
    ws.row_dimensions[r].height = 19
    put(ws, r, 2, None, i + 1)
    for c in range(3, 8):
        put(ws, r, c, None, None, "input")
    put(ws, r, 8, None, None, "input")
    put(ws, r, 9, 10, None, "input", al="left")
BM_LAST = r
dd(ws, "H%d:H%d" % (BM_FIRST, BM_LAST), ["Pass", "Fail", "N/A"])
cf(ws, "H%d:H%d" % (BM_FIRST, BM_LAST))
r += 2
put(ws, r, 2, 10, "BATTERY PROTECTION & BMS SETTINGS", "sec")
r += 1
hdr_spans(ws, r, [(2, 2, "No."), (3, 3, "Parameter"), (4, 4, "Set / Measured"), (5, 5, "Standard / Ref"),
                  (6, 6, "Result"), (7, 10, "Remarks")], h=28)
BP_FIRST = r + 1
for i, p in enumerate(["Over Voltage Protection (V)", "Under Voltage Protection (V)", "Over Current Protection (A)",
                       "Over Temperature Protection (°C)", "Under Temperature Protection (°C)",
                       "Cell Balancing Function Test", "SOC Display Accuracy", "BMS Communication / Monitoring Test",
                       "Battery Emergency Shutdown Test"], 1):
    r += 1
    ws.row_dimensions[r].height = 24
    put(ws, r, 2, None, i)
    put(ws, r, 3, None, p, "plain", al="left")
    put(ws, r, 4, None, None, "input", al="left")
    put(ws, r, 5, None, None, "input", al="left")
    put(ws, r, 6, None, None, "input")
    put(ws, r, 7, 10, None, "input", al="left")
BP_LAST = r
dd(ws, "F%d:F%d" % (BP_FIRST, BP_LAST), ["Pass", "Fail", "N/A"])
cf(ws, "F%d:F%d" % (BP_FIRST, BP_LAST))
r += 2
put(ws, r, 2, 10, "CHARGE / DISCHARGE / BACKUP TEST", "sec")
r += 1
hdr_spans(ws, r, [(2, 3, "Test"), (4, 4, "Start Time"), (5, 5, "End Time"), (6, 6, "Start / End SOC (%)"),
                  (7, 7, "Start / End Voltage (V)"), (8, 8, "Current (A)"), (9, 9, "Duration (min)"), (10, 10, "Remarks")], h=34)
BT_FIRST = r + 1
for t in ["Charge Test", "Discharge Test", "Backup / Off-Grid Test (grid failure simulation)"]:
    r += 1
    ws.row_dimensions[r].height = 28
    put(ws, r, 2, 3, t, "plain", al="left")
    for c in range(4, 10):
        put(ws, r, c, None, None, "input")
    put(ws, r, 10, None, None, "input", al="left")
r += 2
sign_row(ws, r, 10, 5)
MAP["bat"] = {"sheet": BAT, "brand": "D7", "type": "D8", "kwh": "D9", "vn": "D10", "mods": "H7", "strs": "H8", "bms": "H9", "by": "H10", "officer": "D%d" % r,
              "mod": {"first": BM_FIRST, "last": BM_LAST, "cols": {"mod": "C", "str": "D", "vn": "E", "vm": "F", "ins": "G", "s": "H", "rem": "I"}, "num": ["vn", "vm", "ins"]},
              "prm": {"first": BP_FIRST, "cols": {"val": "D", "std": "E", "s": "F", "rem": "G"}},
              "tests": {"first": BT_FIRST, "cols": {"t0": "D", "t1": "E", "soc": "F", "volt": "G", "cur": "H", "dur": "I", "rem": "J"}, "time": ["t0", "t1"]}}
finish(ws)

# =====================================================================
# PROTECTION AND GRID
# =====================================================================
PRO = "Protection and Grid"
ws = new(PRO, {"A": 2.5, "B": 6, "C": 44, "D": 22, "E": 24, "F": 12, "G": 32}, "5B2C6F")
top(ws, 7, "PROTECTION & GRID SETTINGS VERIFICATION")
hdr(ws, 6, ["No.", "Parameter", "Set / Measured Value", "Standard / Reference", "Result", "Remarks"], h=28)
params = ["Over Voltage Protection Trip (V)", "Under Voltage Protection Trip (V)", "Over Frequency Protection Trip (Hz)",
          "Under Frequency Protection Trip (Hz)", "Anti-Islanding Protection Test", "Reconnection Time Delay (s)",
          "Export Power Limit Setting (kW / %)", "RCD / ELCB Rated Trip Current (mA)", "RCD / ELCB Trip Test Result",
          "Earth Fault Protection Test", "Remote Shutdown / Emergency Stop Test",
          "Communication Module (WiFi / 4G / RS485) Test", "Monitoring Portal Registration",
          "Inverter Firmware Version(s)"]
r = 6
PR_FIRST = 7
for i, p in enumerate(params, 1):
    r += 1
    ws.row_dimensions[r].height = 24
    put(ws, r, 2, None, i)
    put(ws, r, 3, None, p, "plain", al="left")
    put(ws, r, 4, None, None, "input", al="left")
    put(ws, r, 5, None, None, "input", al="left")
    put(ws, r, 6, None, None, "input")
    put(ws, r, 7, None, None, "input", al="left")
PR_LAST = r
dd(ws, "F%d:F%d" % (PR_FIRST, PR_LAST), ["Pass", "Fail", "N/A"])
cf(ws, "F%d:F%d" % (PR_FIRST, PR_LAST))
r += 2
put(ws, r, 2, 3, "Verified By", "label")
put(ws, r, 4, 4, None, "input", al="left")
put(ws, r, 5, None, "Date", "label")
put(ws, r, 6, 7, None, "input", fmt="dd-mmm-yyyy")
MAP["prot"] = {"sheet": PRO, "first": PR_FIRST, "cols": {"val": "D", "std": "E", "s": "F", "rem": "G"}, "by": "D%d" % r, "date": "F%d" % r}
finish(ws, landscape=False)

# =====================================================================
# PUNCH LIST
# =====================================================================
PUN = "Punch List"
ws = new(PUN, {"A": 2.5, "B": 5, "C": 12, "D": 18, "E": 34, "F": 14, "G": 14, "H": 30, "I": 16, "J": 12,
               "K": 12, "L": 12}, "5B2C6F")
top(ws, 12, "PUNCH LIST / SNAG LIST")
hdr(ws, 6, ["No.", "Date Raised", "Location / Inverter", "Issue Description", "Category", "Raised By",
            "Corrective Action", "Responsible", "Target Date", "Status", "Closed Date"], h=32)
PU_FIRST = 7
PU_N = 30
for i in range(PU_N):
    r = PU_FIRST + i
    ws.row_dimensions[r].height = 24
    put(ws, r, 2, None, i + 1)
    for c in range(3, 13):
        put(ws, r, c, None, "Open" if False else None, "input", al="left" if c in (4, 5, 7, 8, 9) else None,
            fmt="dd-mmm-yy" if c in (3, 10, 12) else None)
PU_LAST = PU_FIRST + PU_N - 1
MAP["punch"] = {"sheet": PUN, "first": PU_FIRST, "last": PU_LAST, "cols": {"date": "C", "loc": "D", "issue": "E", "cat": "F", "by": "G", "action": "H", "resp": "I", "target": "J", "s": "K", "closed": "L"}, "date": ["date", "target", "closed"]}
dd(ws, "F%d:F%d" % (PU_FIRST, PU_LAST), ["Electrical", "Mechanical", "Civil", "Documentation", "Other"])
dd(ws, "K%d:K%d" % (PU_FIRST, PU_LAST), ["Open", "In Progress", "Closed"])
finish(ws)

# =====================================================================
# SUMMARY (auto)
# =====================================================================
SUM = "Summary"
ws = new(SUM, {"A": 2.5, "B": 5, "C": 40, "D": 12, "E": 9, "F": 9, "G": 9, "H": 16, "I": 36}, "375623")
top(ws, 9, "COMMISSIONING SUMMARY  (calculated automatically)")
ws.row_dimensions[6].height = 22
put(ws, 6, 2, 3, "Project", "label")
put(ws, 6, 4, 9, '=IF(Cover!$C$7="","",Cover!$C$7)', "calc", al="left")
ws.row_dimensions[7].height = 22
put(ws, 7, 2, 3, "Commissioning Date", "label")
put(ws, 7, 4, 9, '=IF(Cover!$F$9="","",Cover!$F$9)', "calc", al="left", fmt="dd-mmm-yyyy")
ws.row_dimensions[8].height = 6
hdr(ws, 9, ["No.", "Section", "Applicable", "Pass", "Fail", "Open", "Status", "Remarks"], h=26)


def cnt(sheet, rng, key):
    return 'COUNTIF(\'%s\'!%s,"%s")' % (sheet, rng, key)


rows = []
rows.append(("Visual & Mechanical Inspection", "Yes",
             cnt(INS, "$D$%d:$D$%d" % (INS_FIRST, INS_LAST), "Pass"),
             cnt(INS, "$D$%d:$D$%d" % (INS_FIRST, INS_LAST), "Fail"),
             "COUNTA('%s'!$C$%d:$C$%d)-%s-%s-%s" % (INS, INS_FIRST, INS_LAST,
                                                   cnt(INS, "$D$%d:$D$%d" % (INS_FIRST, INS_LAST), "Pass"),
                                                   cnt(INS, "$D$%d:$D$%d" % (INS_FIRST, INS_LAST), "Fail"),
                                                   cnt(INS, "$D$%d:$D$%d" % (INS_FIRST, INS_LAST), "N/A")), None))
for _k in range(1, N_INV + 1):
    inv = "INV %02d" % _k
    nm = DCS[inv]
    rng = "$N$%d:$N$%d" % (DC_FIRST, DC_LAST)
    rows.append(("DC String Tests — " + inv, "Yes" if _k <= N_INV_VISIBLE else "No", cnt(nm, rng, "Pass"), cnt(nm, rng, "Fail"),
                 cnt(nm, rng, "Incomplete"), "dc%d" % _k))
rng = "$N$%d:$N$%d" % (A_FIRST, A_LAST)
rows.append(("AC Output (grid no-load & inverters)", "Yes", cnt(ACO, rng, "Pass"), cnt(ACO, rng, "Fail"),
             cnt(ACO, rng, "Incomplete"), None))
rng = "$N$%d:$N$%d" % (I_FIRST, I_LAST)
rows.append(("AC Insulation Resistance", "Yes", cnt(ACI, rng, "Pass"), cnt(ACI, rng, "Fail"),
             cnt(ACI, rng, "Incomplete"), None))
rng = "$G$%d:$G$%d" % (E_FIRST, E_LAST)
rows.append(("Earthing", "Yes", cnt(ACI, rng, "Pass"), cnt(ACI, rng, "Fail"), "0", None))
rows.append(("Thermal Imaging (panels, inverters, switchgear)", "Yes",
             "+".join(cnt(TH, x, "Pass") for x in TH_RES), "+".join(cnt(TH, x, "Fail") for x in TH_RES),
             "+".join(cnt(TH, x, "Incomplete") for x in TH_RES), None))
bm = "$H$%d:$H$%d" % (BM_FIRST, BM_LAST)
bp = "$F$%d:$F$%d" % (BP_FIRST, BP_LAST)
_bmrow = lambda c: "'%s'!$%s$%d:$%s$%d" % (BAT, c, BM_FIRST, c, BM_LAST)
_bm_used = "((%s<>\"\")+(%s<>\"\")+(%s<>\"\")>0)" % (_bmrow("C"), _bmrow("F"), _bmrow("G"))
rows.append(("Battery Test (hybrid systems only)", "No",
             cnt(BAT, bm, "Pass") + "+" + cnt(BAT, bp, "Pass"), cnt(BAT, bm, "Fail") + "+" + cnt(BAT, bp, "Fail"),
             "COUNTA('%s'!$C$%d:$C$%d)-%s-%s-%s+SUMPRODUCT(%s*(%s=\"\"))" % (BAT, BP_FIRST, BP_LAST, cnt(BAT, bp, "Pass"), cnt(BAT, bp, "Fail"),
                                                  cnt(BAT, bp, "N/A"), _bm_used, _bmrow("H")), None))
pr = "$F$%d:$F$%d" % (PR_FIRST, PR_LAST)
rows.append(("Protection & Grid Settings", "Yes", cnt(PRO, pr, "Pass"), cnt(PRO, pr, "Fail"),
             "COUNTA('%s'!$C$%d:$C$%d)-%s-%s-%s" % (PRO, PR_FIRST, PR_LAST, cnt(PRO, pr, "Pass"), cnt(PRO, pr, "Fail"),
                                                   cnt(PRO, pr, "N/A")), None))
pu = "$K$%d:$K$%d" % (PU_FIRST, PU_LAST)
rows.append(("Punch List — all items closed", "Yes", cnt(PUN, pu, "Closed"), "0",
             cnt(PUN, pu, "Open") + "+" + cnt(PUN, pu, "In Progress"), "punch"))
S_FIRST = 10
# "touched" = how many items carry any entry: nothing entered -> NOT TESTED; only N/A entries -> N/A (same rule as the app)
_ins_rng = "'%s'!$D$%d:$D$%d" % (INS, INS_FIRST, INS_LAST)
TOUCHED = {"Visual & Mechanical Inspection": "COUNTA(%s)" % _ins_rng,
           "Protection & Grid Settings": "COUNTA('%s'!$F$%d:$F$%d)" % (PRO, PR_FIRST, PR_LAST),
           "Battery Test (hybrid systems only)": "COUNTA('%s'!$F$%d:$F$%d)+SUMPRODUCT(--(((%s<>\"\")+(%s<>\"\")+(%s<>\"\")+(%s<>\"\"))>0))" % (BAT, BP_FIRST, BP_LAST, _bmrow("C"), _bmrow("F"), _bmrow("G"), _bmrow("H"))}
for i, (name, app, p, f, o, special) in enumerate(rows):
    r = S_FIRST + i
    ws.row_dimensions[r].height = 24
    put(ws, r, 2, None, "=SUBTOTAL(103,$C$%d:C%d)" % (S_FIRST, r))   # numbering skips hidden rows
    put(ws, r, 3, None, name, "plain", al="left")
    put(ws, r, 4, None, app, "input")
    put(ws, r, 5, None, '=IF($D{r}="No",0,{x})'.format(r=r, x=p), "calc")
    put(ws, r, 6, None, '=IF($D{r}="No",0,{x})'.format(r=r, x=f), "calc")
    put(ws, r, 7, None, '=IF($D{r}="No",0,{x})'.format(r=r, x=o), "calc")
    touched = TOUCHED.get(name, "E{r}+F{r}+G{r}".format(r=r))
    if special == "punch":
        stat = '=IF(D{r}="No","N/A",IF(G{r}>0,"IN PROGRESS","PASS"))'.format(r=r)
    else:
        stat = '=IF(D{r}="No","N/A",IF(F{r}>0,"FAIL",IF(({t})=0,"NOT TESTED",IF(G{r}>0,"IN PROGRESS",IF(E{r}>0,"PASS","N/A")))))'.format(r=r, t=touched)
    put(ws, r, 8, None, stat, "calc")
    put(ws, r, 9, None, None, "input", al="left")
    _key = {"Visual & Mechanical Inspection": "visual", "AC Output (grid no-load & inverters)": "acout", "AC Insulation Resistance": "acins",
            "Earthing": "earth", "Thermal Imaging (panels, inverters, switchgear)": "thermal", "Battery Test (hybrid systems only)": "battery",
            "Protection & Grid Settings": "prot", "Punch List — all items closed": "punch"}.get(name)
    if special and special.startswith("dc"):
        _key = special
        if int(special[2:]) > N_INV_VISIBLE:
            ws.row_dimensions[r].hidden = True
    MAP.setdefault("summary", {"sheet": SUM, "col": "D", "rows": {}})["rows"][_key] = r
S_LAST = S_FIRST + len(rows) - 1
dd(ws, "D%d:D%d" % (S_FIRST, S_LAST), ["Yes", "No"])
cf(ws, "H%d:H%d" % (S_FIRST, S_LAST))
OVR = S_LAST + 2
ws.row_dimensions[OVR].height = 30
put(ws, OVR, 2, 7, "OVERALL RESULT", "label")
put(ws, OVR, 8, 9, '=IF(COUNTIF(H{a}:H{b},"FAIL")>0,"FAIL",IF(COUNTIF(H{a}:H{b},"IN PROGRESS")+COUNTIF(H{a}:H{b},"NOT TESTED")>0,"INCOMPLETE","PASS"))'.format(a=S_FIRST, b=S_LAST), "calc")
cf(ws, "H%d:H%d" % (OVR, OVR))
note(ws, OVR + 2, 2, 9,
     "More than two inverters? Unhide the sheet “DC INV 03” (… up to “DC INV 20”), then unhide its row above (select rows %d–%d → right-click → Unhide) and set Applicable to Yes." % (S_FIRST, S_LAST), 28)
finish(ws, landscape=False, one_page=True)

# =====================================================================
# SIGN-OFF
# =====================================================================
SO = "Sign-Off"
ws = new(SO, {"A": 2.5, "B": 30, "C": 24, "D": 24, "E": 24, "F": 20}, "375623")
top(ws, 6, "FINAL COMMISSIONING SIGN-OFF")
ws.row_dimensions[6].height = 26
put(ws, 6, 2, None, "Overall Commissioning Result", "label")
put(ws, 6, 3, 4, "='%s'!$H$%d" % (SUM, OVR), "calc")
cf(ws, "C6:C6")
ws.row_dimensions[7].height = 26
put(ws, 7, 2, None, "Commissioned Date", "label")
put(ws, 7, 3, 4, None, "input", al="left", fmt="dd-mmm-yyyy")
put(ws, 8, 2, None, "Commissioned Officer Name", "label")
ws.row_dimensions[8].height = 26
put(ws, 8, 3, 4, None, "input", al="left")
put(ws, 8, 5, None, "Signature", "label")
put(ws, 8, 6, None, None, "input")
MAP["sign"] = {"sheet": SO, "date": "C7", "officer": "C8", "roles": {"first": 12, "cols": {"name": "C", "co": "D", "date": "F"}}}
put(ws, 10, 2, 6, "SIGN-OFF", "sec")
hdr(ws, 11, ["Role", "Name", "Company", "Signature", "Date"], h=26)
r = 11
for role in ["Commissioning Engineer", "QA / QC Engineer", "Client Representative / Consultant", "Site Safety Officer"]:
    r += 1
    ws.row_dimensions[r].height = 32
    put(ws, r, 2, None, role, "plain", al="left")
    for c in range(3, 7):
        put(ws, r, c, None, None, "input", al="left" if c < 5 else None)
r += 2
put(ws, r, 2, 6, "HANDOVER CHECKLIST", "sec")
r += 1
hdr_spans(ws, r, [(2, 4, "Item"), (5, 6, "Status")], h=24)
h_first = r + 1
for it in ["As-built drawings provided", "O&M manual provided", "Warranty certificates provided",
           "Spare parts list provided", "O&M training conducted", "Monitoring portal access provided",
           "Test certificates / thermal images attached"]:
    r += 1
    ws.row_dimensions[r].height = 22
    put(ws, r, 2, 4, it, "plain", al="left")
    put(ws, r, 5, 6, None, "input")
dd(ws, "E%d:E%d" % (h_first, r), ["Yes", "No", "N/A"])
MAP["sign"]["hand"] = {"first": h_first, "col": "E"}
r += 2
put(ws, r, 2, 6, "COMMENTS", "sec")
ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 5, end_column=6)
MAP["sign"]["comments"] = "B%d" % (r + 1)
ws.cell(r + 1, 2).alignment = Alignment(horizontal="left", vertical="top", wrap_text=True, indent=1)
for rr in range(r + 1, r + 6):
    for c in range(2, 7):
        ws.cell(rr, c).border = BOX
        ws.cell(rr, c).fill = fl(YEL)
finish(ws, landscape=False, one_page=True)

# =====================================================================
# PANEL SERIAL NUMBERS
# =====================================================================
SER = "Panel Serials"
ws = new(SER, {"A": 2.5, "B": 7, "C": 12, "D": 10, "E": 10, "F": 30, "G": 34}, "1F4E78")
top(ws, 7, "PANEL SERIAL NUMBER REGISTER")
hdr(ws, 6, ["No.", "Inverter", "MPPT", "String", "Panel Serial No.", "Remarks"], h=26)
SE_FIRST, SE_N = 7, 60
for i in range(SE_N):
    r = SE_FIRST + i
    ws.row_dimensions[r].height = 18
    put(ws, r, 2, None, i + 1)
    put(ws, r, 3, None, None, "input")
    put(ws, r, 4, None, None, "input")
    put(ws, r, 5, None, None, "input")
    put(ws, r, 6, None, None, "input", al="left")
    put(ws, r, 7, None, None, "input", al="left")
MAP["serials"] = {"sheet": SER, "first": SE_FIRST, "last": SE_FIRST + SE_N - 1, "cols": {"no": "B", "inv": "C", "mppt": "D", "str": "E", "sn": "F", "rem": "G"}}
ws.print_title_rows = "6:6"
finish(ws, landscape=False)

# hidden sheet the app uses to recognise its own exports
APP = "_app"
ws = wb.create_sheet(APP)
ws["A1"] = "SCA-TEMPLATE"; ws["B1"] = TEMPLATE_VERSION
ws["A2"] = "project_id"; ws["A3"] = "exported"; ws["A4"] = "app"; ws["A5"] = "engineer"
ws.sheet_state = "veryHidden"
MAP["app"] = {"sheet": APP, "pid": "B2", "exported": "B3", "app": "B4", "engineer": "B5"}
MAP["dcSummaryKeyPrefix"] = "dc"

# ---- order + save
order = ["Cover", SUM, EQ, INS, INV] + [DCS["INV %02d" % k] for k in range(1, N_INV + 1)] + [SER, ACO, ACI, TH, BAT, PRO, PUN, SO, APP]
wb._sheets = [wb[n] for n in order]
wb.active = 0
wb.save(OUT)
shutil.copy(OUT, RAW)          # un-recalculated copy, embedded in the app as the export template
json.dump(MAP, open(MAPOUT, "w"), indent=1)
print("saved", OUT, len(wb.sheetnames), "sheets")
